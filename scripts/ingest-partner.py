"""Partner-file adapter (spec FR-1.6): a CSV or XLSX of innovations -> .local/pipeline/sources/inn-partner-<slug>.json.

A partner (ROPS, 3 October 2026, "blisko 200" innovations; spec A-05, OP-17) may hand over a spreadsheet
instead of pages. This script turns every row into a source record of the same shape as
scripts/parse-catalogues.py writes (docs/innovation-record.md section 3, schemas/source-record.schema.json),
so that the extraction skill (.claude/skills/extract-innovations) and `derive-records.py build` run unchanged.
Deterministic, no model, no network. A mapping file (YAML) names the columns; a new file layout needs a new
mapping, never a code change (tests/fixtures/partner/mapping.yaml is the commented example).

Shared with the parser (imported from scripts/parse-catalogues.py, not copied): slugify, fingerprint,
build_text and its Polish labels, write_record (a re-run rewrites nothing when only parsed_at would change),
classify_authors, programme_for, material_type, real_urls, PARSER_VERSION and the redaction of parse-v4
(redact_addresses: street addresses and postal codes removed with the town kept, e-mail addresses and phone
numbers replaced by a pointer to the source), applied to every text cell but the title and the links; a row
with such details gets contact_in_source and the review flag redacted:<kind>. A contact column is never
copied. Titles are compared with the build's normalise_title (scripts/derive-records.py).

Ids: inn-partner-<slug>, the slug from the mapped id column if any, else from the title; stable across runs.
Two rows with the same slug get a suffix from a hash of title, organisation and place. A row whose title
equals the title of an existing record (normalised as by the build, FR-1.4) is reported and flagged
possible-duplicate:<id>; the build then lists the pair in duplicates.json for a person to decide.

Usage (from the repository root, with the project venv):
  .venv/Scripts/python scripts/ingest-partner.py --file <csv|xlsx> --map <mapping.yaml> [--sheet NAME] [--root DIR] [--dry-run]
Exit code 1 when the mapping does not fit the file, a row has no title, or a record fails the schema
(valid records are still written, invalid ones are not).
"""
import argparse, csv, datetime, hashlib, importlib.util, io, json, os, re, sys

try:
    import jsonschema, yaml
except ImportError:
    sys.exit("jsonschema or PyYAML is missing: .venv/Scripts/python -m pip install -r requirements.txt")

HERE = os.path.dirname(os.path.abspath(__file__))


def load_module(name, filename):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, filename))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


parser = load_module("parse_catalogues", "parse-catalogues.py")
referee = load_module("derive_records", "derive-records.py")

COLUMN_ROLES = ("id", "title", "intro", "url", "licence", "category", "implementer", "organisation",
                "organisation_website", "authors", "place", "incubator", "incubator_years",
                "dissemination_label", "contact", "materials")
NOT_REDACTED = {"url", "organisation_website", "materials", "contact", "licence"}
SLUG_MAX = 70


def log(msg):
    print(msg, flush=True)


def fold(s):
    return referee.normalise_title(s)


# ---------------------------------------------------------------- reading the file

def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (datetime.datetime, datetime.date)):
        v = v.isoformat()[:10]
    s = str(v).replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    return "\n".join(line.rstrip() for line in s.split("\n")).strip()


def read_rows(path, sheet):
    """All rows of the file as lists of strings, with their 1-based row number in the file."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            sys.exit("openpyxl is missing: .venv/Scripts/python -m pip install -r requirements.txt")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet and sheet not in wb.sheetnames:
            sys.exit(f"sheet {sheet!r} not in {wb.sheetnames}")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        return [(i, [cell_text(v) for v in row]) for i, row in enumerate(ws.iter_rows(values_only=True), 1)]
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp1250"):   # Polish Excel saves CSV as Windows-1250
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    body = "\n".join(l for l in text.splitlines() if not l.startswith("#"))[:8192]
    try:
        dialect = csv.Sniffer().sniff(body, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    out, lineno = [], 1
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    for row in reader:
        out.append((lineno, [cell_text(v) for v in row]))
        lineno = reader.line_num + 1
    return out


def norm_header(h):
    return re.sub(r"\s+", " ", h or "").strip().lower()


def find_header(rows, title_column, header_row):
    if header_row:
        for n, (lineno, cells) in enumerate(rows):
            if lineno == header_row:
                return n
        sys.exit(f"header_row {header_row} is past the end of the file")
    for n, (lineno, cells) in enumerate(rows):
        if cells and cells[0].startswith("#"):
            continue
        if norm_header(title_column) in [norm_header(c) for c in cells]:
            return n
    sys.exit(f"no row carries the title column {title_column!r}; set file.header_row in the mapping")


# ---------------------------------------------------------------- redaction and mapping

def redact(text):
    """The parser's redaction (parse-v4, redact_addresses): street addresses and postal codes removed with the
    town kept, e-mail addresses and phone numbers replaced by a pointer to the source. Returns the kinds found."""
    if not text:
        return text, []
    out = parser.redact_addresses(text)
    if out == text:
        return out, []
    kinds = [k for k, rx in (("address", parser.ADDRESS_RE), ("email", parser.EMAIL_RE)) if rx.search(text)]
    if any(len(re.sub(r"\D", "", m.group(0))) in (9, 11) for m in parser.PHONE_RE.finditer(text)):
        kinds.append("phone")
    return out, kinds or ["other"]


def split_list(value, sep):
    return [p.strip() for p in re.split("|".join(re.escape(s) for s in sep), value) if p.strip()] if value else []


def map_codes(values, items, keys):
    """Codes whose code or one of the named label keys equals the value (case and diacritics folded)."""
    codes, unknown = [], []
    for v in values:
        hit = [t["code"] for t in items if any(t.get(k) and fold(t[k]) == fold(v) for k in ("code",) + keys)]
        if hit:
            codes += hit
        else:
            unknown.append(v)
    return sorted(set(codes)), unknown


def make_slug(value):
    s = parser.slugify(value)
    if len(s) > SLUG_MAX:
        s = s[:SLUG_MAX].rsplit("-", 1)[0]
    return s or "bez-tytulu"


# ---------------------------------------------------------------- main

def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Partner-file adapter (FR-1.6): CSV or XLSX -> source records")
    ap.add_argument("--file", required=True)
    ap.add_argument("--map", required=True)
    ap.add_argument("--sheet")
    ap.add_argument("--root", default=os.path.dirname(HERE), help="repository to write into (default: this one)")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    root = os.path.abspath(args.root)
    out_dir = os.path.join(root, ".local", "pipeline", "sources")
    mapping = yaml.safe_load(open(args.map, encoding="utf-8"))
    src_cfg, file_cfg = mapping.get("source") or {}, mapping.get("file") or {}
    columns = {k: v for k, v in (mapping.get("columns") or {}).items() if v}
    fields_map = mapping.get("source_fields") or {}
    unknown_roles = sorted(set(columns) - set(COLUMN_ROLES))
    if unknown_roles:
        sys.exit(f"mapping: unknown column roles {unknown_roles}; known: {list(COLUMN_ROLES)}")
    if "title" not in columns or not fields_map:
        sys.exit("mapping: columns.title and at least one source_fields entry are required")
    for key, label in (mapping.get("labels") or {}).items():
        parser.LABELS_PL[key] = label          # labels for text_pl, in this process only
    sep = file_cfg.get("list_separator") or [";", "\n"]
    sep = [sep] if isinstance(sep, str) else sep
    required = mapping.get("required") or []
    source_name = src_cfg.get("name") or "partner-rops"
    retrieved_at = str(src_cfg.get("retrieved_at") or datetime.date.today().isoformat())
    licence = src_cfg.get("licence") or parser.CC_BY
    licence_url = src_cfg.get("licence_url") or (parser.CC_BY_URL if licence == parser.CC_BY else "")
    order = list(fields_map)

    raw = open(args.file, "rb").read()
    raw_sha = hashlib.sha256(raw).hexdigest()
    in_path = os.path.abspath(args.file)
    raw_path = os.path.relpath(in_path, root).replace("\\", "/") if in_path.lower().startswith(root.lower() + os.sep) else in_path.replace("\\", "/")

    rows = read_rows(args.file, args.sheet or file_cfg.get("sheet"))
    h = find_header(rows, columns["title"], file_cfg.get("header_row"))
    header = [norm_header(c) for c in rows[h][1]]

    def col_index(name):
        n = norm_header(name)
        return header.index(n) if n in header else None

    mat_cols = columns.get("materials") or []
    mat_cols = [mat_cols] if isinstance(mat_cols, str) else mat_cols
    wanted = {**{f"columns.{r}": c for r, c in columns.items() if r != "materials"},
              **{f"columns.materials[{i}]": c for i, c in enumerate(mat_cols)},
              **{f"source_fields.{k}": c for k, c in fields_map.items()}}
    missing_cols = {k: c for k, c in wanted.items() if col_index(c) is None}
    if missing_cols:
        log("the mapping names columns the file does not have:")
        for k, c in missing_cols.items():
            log(f"  {k}: {c!r}")
        log("columns of the file: " + ", ".join(repr(c) for c in rows[h][1] if c))
        return 1

    tax = json.load(open(os.path.join(root, "data", "taxonomies.json"), encoding="utf-8"))
    schema = jsonschema.Draft202012Validator(json.load(open(os.path.join(root, "schemas", "source-record.schema.json"), encoding="utf-8")))
    existing = {}
    if os.path.isdir(out_dir):
        for fn in sorted(os.listdir(out_dir)):
            if fn.endswith(".json") and not fn.startswith("inn-partner-"):
                rec = json.load(open(os.path.join(out_dir, fn), encoding="utf-8"))
                existing.setdefault(fold(rec["title"]), []).append(rec["id"])
    existing_slugs = {fn[:-5].split("-", 2)[2]: fn[:-5] for fn in os.listdir(out_dir) if fn.endswith(".json")} if os.path.isdir(out_dir) else {}

    errors, records, reports = 0, [], []
    for lineno, cells in rows[h + 1:]:
        if not any(cells) or (cells and cells[0].startswith("#")):
            continue

        def get(col):
            i = col_index(col)
            return cells[i] if i is not None and i < len(cells) else ""

        vals = {r: get(c) for r, c in columns.items() if r != "materials"}
        title = vals.get("title", "")
        if not title:
            log(f"row {lineno}: no title, skipped")
            errors += 1
            continue
        flags, redacted = [f"partner-row:{lineno}"], set()

        fields = {}
        for key, col in fields_map.items():
            v = get(col)
            if v:
                v, kinds = redact(v)
                redacted.update(kinds)
                if v:
                    fields[key] = v
        for r in list(vals):
            if r not in NOT_REDACTED and r != "title" and vals[r]:
                vals[r], kinds = redact(vals[r])
                redacted.update(kinds)
        for key in required:
            if key not in fields:
                flags.append(f"missing-field:{key}")
        flags += [f"redacted:{k}" for k in sorted(redacted)]

        # target groups from the category column: rops_category_slug, code or label of data/taxonomies.json
        cats = split_list(vals.get("category"), sep)
        tg, unknown = map_codes(cats, tax["target_groups"], ("rops_category_slug", "label_pl"))
        flags += [f"unknown-category:{u}" for u in unknown]
        if cats and not tg:
            tg = ["inne"]
        if not cats:
            flags.append("missing-field:category")
        cat_slug = next((t["rops_category_slug"] for t in tax["target_groups"] if tg and t["code"] == tg[0]), None)
        category = cat_slug or (parser.slugify(cats[0]) if cats else None)
        it, unknown = map_codes(split_list(vals.get("implementer"), sep), tax["implementer_types"], ("label_pl", "national_base_value"))
        flags += [f"unknown-implementer:{u}" for u in unknown]

        # organisation, persons: the organisation column first, then the parser's reading of the authors
        orgs, persons, unclassified = parser.classify_authors("\n".join(split_list(vals.get("authors"), sep)))
        if unclassified:
            flags.append("authors-unclassified")
        org_name = vals.get("organisation") or (orgs[0] if orgs else None)
        others = [o for o in orgs if o != org_name]
        website = next(iter(parser.real_urls(vals.get("organisation_website", ""))), None) if org_name else None
        organisation = {"name": org_name, "website": website, **({"others": others} if others else {})} if org_name else None
        links = [{"title": "Strona internetowa", "url": website}] if website else []

        url_cell = parser.real_urls(vals.get("url", "")) if vals.get("url") else []
        entry_url = next((u for u in url_cell if u.startswith("https://")), None)
        links += [{"title": "Link", "url": u} for u in url_cell if u != entry_url]
        row_licence = vals.get("licence") or licence
        materials = []
        for col in mat_cols:
            for u in parser.real_urls(get(col)):
                mt = parser.material_type(u)
                materials.append({"type": mt, "title": {"pdf": "Folder informacyjny", "video": "Film", "zip": "Materiały do pobrania"}.get(mt, "Link"),
                                  "url": u, "licence": None if mt == "video" else row_licence, "local_path": None})

        incubator = vals.get("incubator") or None
        rops_inc = bool(incubator) and any(n in incubator.lower() for n in parser.ROPS_INCUBATORS)
        label = vals.get("dissemination_label") or None
        origin = {
            "incubator_name": incubator, "incubator_years": vals.get("incubator_years") or None, "incubator_profile_url": None,
            "programme": parser.programme_for(vals.get("incubator_years")) if incubator else None,
            "selected_for_dissemination": True if label else None, "dissemination_label_pl": label,
            "region": "małopolskie" if rops_inc else None, "rops_incubated": rops_inc,
        }
        intro = vals.get("intro") or None
        records.append({"lineno": lineno, "slug_base": make_slug(vals.get("id") or title),
                        "hash": hashlib.sha256(json.dumps([title, org_name, vals.get("place")], ensure_ascii=False).encode("utf-8")).hexdigest()[:6],
                        "rec": {
            "id": None, "source": source_name, "title": title, "intro_pl": intro,
            "sources": [{"name": source_name, "url": entry_url, "category_slug": category, "retrieved_at": retrieved_at,
                         "licence": row_licence, "licence_url": licence_url, "raw_path": raw_path, "raw_sha256": raw_sha}],
            "source_fields": fields,
            "tags": {"advanced": [], "simple": []},
            "mapped": {"target_groups": tg, "implementer_types": it, "character": [], "tools": [], "areas": [], "settings": [], "domain_hints": []},
            "innovator": {"type": None, "place_name": vals.get("place") or None},
            "organisation": organisation, "persons_public": persons,
            **({"authors_unclassified": unclassified} if unclassified else {}),
            "contact_in_source": bool(vals.get("contact")) or bool(redacted),
            "origin": origin, "materials": materials, "links": links,
            "text_pl": parser.redact_addresses(parser.build_text(title, intro, fields, order)),
            "review_flags": None,
            "fingerprint": parser.fingerprint(title, category, fields),
            "parser_version": parser.PARSER_VERSION,
            "parsed_at": datetime.datetime.now().isoformat(timespec="seconds"),
        }, "flags": flags})

    # ids: unique and stable; a slug shared by several rows gets the row's hash
    counts = {}
    for r in records:
        counts[r["slug_base"]] = counts.get(r["slug_base"], 0) + 1
    seen = set()
    for r in records:
        slug = r["slug_base"] if counts[r["slug_base"]] == 1 else f"{r['slug_base']}-{r['hash']}"
        if slug in seen:
            log(f"row {r['lineno']}: same title, organisation and place as an earlier row; skipped")
            errors += 1
            r["rec"] = None
            continue
        seen.add(slug)
        rec = r["rec"]
        rec["id"] = f"inn-partner-{slug}"
        if counts[r["slug_base"]] > 1:
            reports.append(f"row {r['lineno']}: {counts[r['slug_base']]} rows share the slug {r['slug_base']!r}; id {rec['id']}")
        dups = existing.get(fold(rec["title"]), [])
        same_slug = existing_slugs.get(slug)
        for d in sorted(set(dups + ([same_slug] if same_slug and not same_slug.startswith("inn-partner-") else []))):
            r["flags"].append(f"possible-duplicate:{d}")
            reports.append(f"row {r['lineno']}: {rec['id']} has the title of the existing record {d} (the build lists the pair, FR-1.4)")
        rec["review_flags"] = sorted(set(r["flags"]))

    written = unchanged = invalid = 0
    produced = set()
    for r in records:
        rec = r["rec"]
        if rec is None:
            continue
        errs = [f"{'/'.join(str(p) for p in e.absolute_path) or '(root)'}: {e.message[:160]}" for e in schema.iter_errors(rec)]
        path = os.path.join(out_dir, rec["id"] + ".json")
        if errs:
            invalid += 1
            errors += 1
            log(f"row {r['lineno']}: {rec['id']} INVALID against schemas/source-record.schema.json, not written")
            for e in errs:
                log(f"  {e}")
            continue
        produced.add(rec["id"])
        state = "new"
        if os.path.exists(path):
            old = json.load(open(path, encoding="utf-8"))
            same = {k: v for k, v in old.items() if k != "parsed_at"} == {k: v for k, v in rec.items() if k != "parsed_at"}
            state = "unchanged" if same else "changed"
        log(f"row {r['lineno']}: {rec['id']} [{state}] tg={rec['mapped']['target_groups']} it={rec['mapped']['implementer_types']} flags={rec['review_flags']}")
        if args.dry_run:
            written += state != "unchanged"
            unchanged += state == "unchanged"
            continue
        if parser.write_record(path, rec):
            written += 1
        else:
            unchanged += 1
    for line in reports:
        log("report: " + line)
    if os.path.isdir(out_dir):
        stale = sorted(fn[:-5] for fn in os.listdir(out_dir) if fn.startswith("inn-partner-") and fn[:-5] not in produced)
        if stale:
            log(f"report: {len(stale)} partner records are not in this file and were left in place: {', '.join(stale)}")
    verb = "would be written" if args.dry_run else "written or changed"
    log(f"{len(produced)} valid records ({written} {verb}, {unchanged} unchanged){' [dry run]' if args.dry_run else ''}; "
        f"{invalid} invalid; {errors} errors; file sha256 {raw_sha[:12]}; into {out_dir}")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
